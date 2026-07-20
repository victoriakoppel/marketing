#!/usr/bin/env python3
# Builds the source-of-truth workbook: edit salaries -> totals recalc automatically.
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# key, dept, block, typ, label(line1), detail(line2), salary, assumed, vacant
ROWS = [
    # Marcom
    ("M_HEAD","Marcom","Руководитель","Руководитель","Marcom · Олег С","Столярчук Олег Валерьевич",3167500,"",""),
    ("M_PROMO_1","Marcom","Промо / Контент","Сотрудник","Senior Promo Original · Аида А","Әлішбай Аида Шералықызы",1247076,"",""),
    ("M_PROMO_2","Marcom","Промо / Контент","Сотрудник","Promo · Александр К","Казачков Александр",800000,"",""),
    ("M_PROMO_3","Marcom","Промо / Контент","Сотрудник","Promo · Райлана","Жанибекова Райлана Нурлановна",1000162,"",""),
    ("M_PROMO_4","Marcom","Промо / Контент","Сотрудник","Promo · Яков","Сорокин Яков Андреевич",614931,"",""),
    ("M_PROMO_5","Marcom","Промо / Контент","Сотрудник","Promo · Чолпон","Эркинова Чолпон",1025642,"",""),
    ("M_SMM_1","Marcom","SMM","Сотрудник","Head of SMM · Мария","Коренькова Мария Николаевна",1625032,"",""),
    ("M_SMM_2","Marcom","SMM","Сотрудник","SMM Original","SMM менеджер · имя н/д",867456,"макс.",""),
    ("M_SMM_3","Marcom","SMM","Сотрудник","SMM Gen","SMM менеджер · имя н/д",867456,"макс.",""),
    ("M_SMM_4","Marcom","SMM","Сотрудник","SMM Sport","SMM менеджер · имя н/д",867456,"макс.",""),
    ("M_SMM_5","Marcom","SMM","Сотрудник","SMM Original","SMM менеджер · имя н/д",867456,"макс.",""),
    ("M_SMM_6","Marcom","SMM","Сотрудник","Telegram редактор Sport","Куричев Иван Николаевич",513920,"",""),
    ("M_PR_1","Marcom","PR","Сотрудник","Pr manager · Аружан","PR менеджер",1212314,"",""),
    # Каналы коммуникации
    ("K_HEAD","Каналы коммуникации","Руководитель","Руководитель","Head · Эдуард Д","Head каналы коммуникации",2260000,"",""),
    ("K_1","Каналы коммуникации","—","Сотрудник","Senior PPC · Антон Б","Senior PPC",1346342,"",""),
    ("K_2","Каналы коммуникации","—","Сотрудник","Media manager · Халмурад","Медиа менеджер",1056850,"",""),
    ("K_3","Каналы коммуникации","—","Сотрудник","Внут. маркетинг · Аида Е","Внутренний маркетинг",1058434,"",""),
    ("K_4","Каналы коммуникации","—","Сотрудник","PPC (1)","Трафик менеджер · имя н/д",1124000,"макс.",""),
    ("K_5","Каналы коммуникации","—","Сотрудник","CVM (?)","CRM менеджер · планируется",867456,"макс.","да"),
    # Продакшн
    ("P_HEAD","Продакшн","Руководитель","Руководитель","Head of creative · Катерина Р","Креативный директор",2490895,"",""),
    ("P_ART","Продакшн","Арт","Суб-лид","Senior art · Олег Н","Назаренко Олег Станиславович",1494981,"",""),
    ("P_PROD","Продакшн","Продакшн","Суб-лид","Senior production · Александр Е","Senior production",1291685,"",""),
    ("P_ART_1","Продакшн","Арт","Сотрудник","Graph","Дизайнер · имя н/д",940000,"макс.",""),
    ("P_ART_2","Продакшн","Арт","Сотрудник","Graph","Дизайнер · имя н/д",940000,"макс.",""),
    ("P_ART_3","Продакшн","Арт","Сотрудник","Graph Спорт","Дизайнер · имя н/д",940000,"макс.",""),
    ("P_ART_4","Продакшн","Арт","Сотрудник","Motion","Motion-дизайнер · имя н/д",877000,"макс.",""),
    ("P_ART_5","Продакшн","Арт","Сотрудник","Motion","Motion-дизайнер · имя н/д",877000,"макс.",""),
    ("P_PROD_1","Продакшн","Продакшн","Сотрудник","Мобилограф / монтажёр","Видеограф · имя н/д",1498769,"макс.",""),
    ("P_PROD_2","Продакшн","Продакшн","Сотрудник","Мобилограф / монтажёр","Видеограф · имя н/д",1498769,"макс.",""),
    ("P_PROD_3","Продакшн","Продакшн","Сотрудник","Мобилограф / монтажёр","Видеограф · имя н/д",1498769,"макс.",""),
    ("P_PROD_4","Продакшн","Продакшн","Сотрудник","Мобилограф / монтажёр Sport","Видеограф · имя н/д",1498769,"макс.",""),
    # Проджект
    ("PJ_HEAD","Проджект","Руководитель","Руководитель","Head · Ольга М","Head project",1291685,"",""),
    ("PJ_1","Проджект","—","Сотрудник","Project manager · Аружан С","Проджект-менеджер",700000,"",""),
    ("PJ_2","Проджект","—","Сотрудник","Project manager · Анель","Муханова Анель Султановна",741192,"",""),
    ("PJ_3","Проджект","—","Сотрудник","Спецпроекты · Марина","Гяурова Марина Якубовна",1000162,"",""),
    ("PJ_4","Проджект","—","Сотрудник","B2B и продакшн · Олеся М","B2B и продакшн",1081000,"",""),
]
DEPTS = ["Marcom","Каналы коммуникации","Продакшн","Проджект"]

NAVY="1F3A5F"; ICE="EAF1FB"; MUTED="5B6B7C"
BLUE_INPUT="0000CC"; AMBER_FILL="FFF3E0"; VAC_FILL="FBE4E4"
thin=Side(style="thin",color="C7D8EE")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

wb = openpyxl.Workbook()
# Force Excel to recalculate all formulas when the file is opened (no cached values from openpyxl).
try:
    wb.calculation.fullCalcOnLoad = True
except Exception:
    pass

# ---------------- Sheet: Данные ----------------
ws = wb.active; ws.title = "Данные"
headers = ["№","Отдел","Блок","Тип","Позиция (строка 1)","Подпись (строка 2)","Оклад, ₸","Признак","Вакансия","Ключ"]
ws.merge_cells("A1:J1")
t = ws["A1"]; t.value = "Данные — команда маркетинга и ФОТ  ·  редактируйте синие ячейки в столбце «Оклад, ₸»"
t.font = Font(name="Arial",bold=True,size=13,color="FFFFFF"); t.alignment=Alignment(vertical="center",horizontal="left",indent=1)
t.fill = PatternFill("solid",fgColor=NAVY); ws.row_dimensions[1].height = 26
for j,h in enumerate(headers,1):
    c = ws.cell(row=2,column=j,value=h)
    c.font=Font(name="Arial",bold=True,size=10,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY)
    c.alignment=Alignment(vertical="center",horizontal="center",wrap_text=True); c.border=border
ws.row_dimensions[2].height=30
r=3
for i,(key,dept,block,typ,label,detail,sal,flag,vac) in enumerate(ROWS,1):
    vals=[i,dept,block,typ,label,detail,sal,flag,vac,key]
    for j,v in enumerate(vals,1):
        c=ws.cell(row=r,column=j,value=v); c.border=border
        c.font=Font(name="Arial",size=10,color=MUTED if j in (3,4,6,8,9,10) else "1F2A37")
        c.alignment=Alignment(vertical="center",horizontal="left",wrap_text=(j in (5,6)))
    # salary cell = editable input (blue)
    sc=ws.cell(row=r,column=7)
    sc.font=Font(name="Arial",size=10,bold=True,color=BLUE_INPUT)
    sc.number_format='#,##0" ₸"'; sc.alignment=Alignment(vertical="center",horizontal="right")
    ws.cell(row=r,column=1).alignment=Alignment(vertical="center",horizontal="center")
    ws.cell(row=r,column=8).alignment=Alignment(vertical="center",horizontal="center")
    ws.cell(row=r,column=9).alignment=Alignment(vertical="center",horizontal="center")
    if flag=="макс.":
        sc.fill=PatternFill("solid",fgColor=AMBER_FILL)
    if vac=="да":
        for j in range(1,11): ws.cell(row=r,column=j).fill=PatternFill("solid",fgColor=VAC_FILL)
    r+=1
last=r-1
widths=[5,20,16,13,30,32,15,10,10,11]
for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes="A3"
# note row
nr=last+2
ws.merge_cells(start_row=nr,start_column=1,end_row=nr,end_column=10)
n=ws.cell(row=nr,column=1,value="Синие ячейки в «Оклад, ₸» — редактируемые. «Признак: макс.» — имя на позиции не определено, подставлена максимальная стоимость по позиции. «Вакансия: да» — планируемая позиция. Итоги считаются на листе «Свод».")
n.font=Font(name="Arial",size=9,italic=True,color=MUTED); n.alignment=Alignment(wrap_text=True,vertical="top")
ws.row_dimensions[nr].height=30

# ---------------- Sheet: Свод ----------------
sv=wb.create_sheet("Свод")
sv.merge_cells("A1:C1")
h=sv["A1"]; h.value="Свод по отделам — численность и ФОТ (пересчитывается автоматически)"
h.font=Font(name="Arial",bold=True,size=13,color="FFFFFF"); h.fill=PatternFill("solid",fgColor=NAVY)
h.alignment=Alignment(vertical="center",indent=1); sv.row_dimensions[1].height=26
for j,cap in enumerate(["Отдел","Численность, чел.","ФОТ, ₸ / мес"],1):
    c=sv.cell(row=2,column=j,value=cap); c.font=Font(name="Arial",bold=True,size=11,color="FFFFFF")
    c.fill=PatternFill("solid",fgColor=NAVY); c.alignment=Alignment(vertical="center",horizontal="center",wrap_text=True); c.border=border
sv.row_dimensions[2].height=28
rng_dept=f"Данные!$B$3:$B${last}"; rng_sal=f"Данные!$G$3:$G${last}"
rr=3
for d in DEPTS:
    a=sv.cell(row=rr,column=1,value=d); a.font=Font(name="Arial",size=11,bold=True,color="1F2A37")
    b=sv.cell(row=rr,column=2,value=f'=COUNTIF({rng_dept},A{rr})'); b.alignment=Alignment(horizontal="center")
    b.font=Font(name="Arial",size=11,color="1F2A37")
    c=sv.cell(row=rr,column=3,value=f'=SUMIF({rng_dept},A{rr},{rng_sal})')
    c.font=Font(name="Arial",size=11,bold=True,color="0F6E4F"); c.number_format='#,##0" ₸"'
    for j in range(1,4):
        cell=sv.cell(row=rr,column=j); cell.border=border
        cell.fill=PatternFill("solid",fgColor=ICE if rr%2 else "FFFFFF")
    rr+=1
# total
a=sv.cell(row=rr,column=1,value="Итого"); a.font=Font(name="Arial",size=12,bold=True,color="FFFFFF"); a.fill=PatternFill("solid",fgColor=NAVY)
b=sv.cell(row=rr,column=2,value=f'=SUM(B3:B{rr-1})'); b.font=Font(name="Arial",size=12,bold=True,color="FFFFFF"); b.fill=PatternFill("solid",fgColor=NAVY); b.alignment=Alignment(horizontal="center")
c=sv.cell(row=rr,column=3,value=f'=SUM(C3:C{rr-1})'); c.font=Font(name="Arial",size=12,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=NAVY); c.number_format='#,##0" ₸"'
for j in range(1,4): sv.cell(row=rr,column=j).border=border
sv.column_dimensions["A"].width=26; sv.column_dimensions["B"].width=18; sv.column_dimensions["C"].width=20
sv.sheet_view.showGridLines=False
nr2=rr+2
sv.merge_cells(start_row=nr2,start_column=1,end_row=nr2,end_column=3)
nn=sv.cell(row=nr2,column=1,value="Числа берутся с листа «Данные». Измените оклад там — итоги здесь обновятся автоматически. Презентация пересобирается из этого файла (см. build/README).")
nn.font=Font(name="Arial",size=9,italic=True,color=MUTED); nn.alignment=Alignment(wrap_text=True,vertical="top")
sv.row_dimensions[nr2].height=42

import os
out=os.environ.get("XLSX_OUT") or os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","Данные — структура и ФОТ.xlsx")
wb.save(out)
print("saved", out, "rows:", len(ROWS))
